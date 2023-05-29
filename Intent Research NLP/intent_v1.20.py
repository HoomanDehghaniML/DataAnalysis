import os
import pandas as pd
from google.cloud import language_v1
from google.cloud.language_v1 import types
import shutil
import openpyxl

def NLP(competitors, sq1, competitor_url, sq1_url):
    # Set up Google Natural Language API
    os.environ['GOOGLE_APPLICATION_CREDENTIALS'] = r"C:\Users\Hooman Deghani\OneDrive\PC\Desktop\API Keys\natural-language-api-381923-52cfb8b44bd7.json"
    client = language_v1.LanguageServiceClient()

    def get_entities(text):
        document = language_v1.Document(content=text, type_=language_v1.Document.Type.PLAIN_TEXT)
        request = language_v1.AnalyzeEntitiesRequest(document=document)
        response = client.analyze_entities(request=request)
        return [(entity.name, entity.salience) for entity in response.entities]

    def get_categories(text):
        document = language_v1.Document(content=text, type_=language_v1.Document.Type.PLAIN_TEXT)
        request = language_v1.ClassifyTextRequest(document=document)
        response = client.classify_text(request=request)
        return [(category.name, category.confidence) for category in response.categories]

    def get_sentiment(text):
        document = language_v1.Document(content=text, type_=language_v1.Document.Type.PLAIN_TEXT)
        request = language_v1.AnalyzeSentimentRequest(document=document)
        response = client.analyze_sentiment(request=request)
        return response.document_sentiment.score

    def get_sentiment_data(text):
        document = language_v1.Document(content=text, type_=language_v1.Document.Type.PLAIN_TEXT)
        request = language_v1.AnalyzeSentimentRequest(document=document)
        response = client.analyze_sentiment(request=request)
        sentiment_data = []
        for sentence in response.sentences:
            sentiment_data.append((sentence.text.content, sentence.sentiment.score, sentence.sentiment.magnitude))
        return sentiment_data

    def read_text_file(file_path):
        with open(file_path, 'r', encoding='utf-8') as file:
            return file.read()

    competitor_data = [read_text_file(file_path) for file_path in competitors]
    sq1_data = read_text_file(sq1)

    # Get entity data
    entity_data = []
    for text in competitor_data:
        entity_data.extend(get_entities(text))

    df_competitors = pd.DataFrame(entity_data, columns=['Entity', 'Salience'])
    df_competitors_aggregated = df_competitors.groupby('Entity').mean().reset_index()

    df_sq1 = pd.DataFrame(get_entities(sq1_data), columns=['Entity', 'Salience'])

    common_entities = df_competitors_aggregated[df_competitors_aggregated['Entity'].isin(df_sq1['Entity'])]
    df_competitors_final = df_competitors_aggregated[~df_competitors_aggregated['Entity'].isin(common_entities['Entity'])]

    # Sort df_competitors_final by Salience score in descending order
    df_competitors_final = df_competitors_final.sort_values('Salience', ascending=False)

    # Remove rows with Salience < 0.00055
    df_competitors_final = df_competitors_final[df_competitors_final['Salience'] >= 0.00055]

    # Save entity data to xlsx
    df_competitors_final.to_excel("entities.xlsx", index=False, engine='openpyxl')

    # Get category data
    competitor_categories = []
    for text in competitor_data:
        competitor_categories.extend(get_categories(text))

    df_category_competitors = pd.DataFrame(competitor_categories, columns=['Category', 'Confidence score'])
    df_category_competitors_aggregated = df_category_competitors.groupby('Category').mean().reset_index()
    df_category_sq1 = pd.DataFrame(get_categories(sq1_data), columns=['Category', 'Confidence score'])

    category_html = f"<h1> Category Data </h1> <h2> Target Categories: </h2>{df_category_competitors_aggregated.to_html()} <h2> Our Categories</h2>{df_category_sq1.to_html()} <p> If the top three categories from both tables are different, then the article needs an <strong>edit</strong>."

    # Get sentiment data
    sentiment_competitors = sum(get_sentiment(text) for text in competitor_data) / len(competitor_data)
    sentiment_sq1 = get_sentiment(sq1_data)
    sentiment_difference = sentiment_competitors - sentiment_sq1
    sentiment_html = sentiment_html = f"<h1>Sentiment Difference</h1><p>Target Sentiment: {sentiment_competitors}</p> <p> Sq1 Sentiment: {sentiment_sq1}</p><p>Sentiment difference is {sentiment_difference}.\n</p>"
    if abs(sentiment_difference) > 0.1:
        sentiment_html += "<p>We need to <strong>edit</strong> the article.</p>"

    # New functionality: Get sentiment data for sq1 and filter based on sentiment_difference
    df_sentiment = pd.DataFrame(get_sentiment_data(sq1_data), columns=['phrase', 'sentiment score', 'magnitude'])
    if sentiment_difference > 0:
        df_sentiment = df_sentiment[df_sentiment['sentiment score'] < 0]
    else:
        df_sentiment = df_sentiment[df_sentiment['sentiment score'] > 0]

    # Save sentimentdDataFrame to Excel file
    df_sentiment.to_excel("sentiment_phrases.xlsx", index=False, engine='openpyxl')

    # Open the workbook
    workbook = openpyxl.load_workbook("sentiment_phrases.xlsx")
    worksheet = workbook.active

    # Adjust the column width to fit the content
    for column_cells in worksheet.columns:
        length = max(len(str(cell.value)) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = length

    # Save the workbook
    workbook.save("sentiment_phrases.xlsx")

    # Create a new folder and save the files
    folder_name = input("Enter the name of the folder: ")
    output_dir = r"G:\Shared drives\sq1 - marketing\seo\Intent Research"
    folder_path = os.path.join(output_dir, folder_name)

    # Save summary to HTML file
    with open('summary.html', 'w') as summary_file:
    # Include links to competitor and sq1 pages
        summary_file.write("<h1>Articles</h1>")
        for i, url in enumerate(competitor_urls):
            summary_file.write(f'<p><a href="{url}">Competitor {i+1}</a></p>\n')
        summary_file.write(f'<p><a href="{sq1_url}">SQ1</a></p>\n\n')

        summary_file.write("<html><head><style>table, th, td {border: 1px solid black; border-collapse: collapse;} th, td {padding: 10px;} </style></head><body>")
        summary_file.write(category_html)
        summary_file.write("<h1>Edit</h1>")
        if sentiment_html:
            summary_file.write(sentiment_html)
        summary_file.write(f"Go through sentiment_phrases.xlsx and adjust the sentences in the sq1 article based on sentiment difference. Make sure to only change the tone of sentences and not change words/phrases that are salient to the text.")
        summary_file.write(f'<p><a href="{folder_path}/sentiment_phrases.xlsx">sentiment_phrases.xlsx</a></p>')
        summary_file.write(f"Go through entities.xlsx and add the entities to the sq1 article in places you find appropriate.")
        summary_file.write(f'<p><a href="{folder_path}/entities.xlsx">entities.xlsx</a></p>') # Add a link to the entities.xlsx file
        summary_file.write("</body></html>")

    # Create the folder if it doesn't exist
    if not os.path.exists(folder_path):
        os.makedirs(folder_path)

    # Save the files in the new folder
    shutil.move("entities.xlsx", os.path.join(folder_path, "entities.xlsx"))
    shutil.move("summary.html", os.path.join(folder_path, "summary.html"))
    shutil.move("sentiment_phrases.xlsx", os.path.join(folder_path, "sentiment_phrases.xlsx"))

    print("Done!")

# User data
competitors_list = [
r"C:\Users\Hooman Deghani\Python\Data Analysis\HTML\https___www.investopedia.com_updates_first-time-home-buyer_.txt",
r"C:\Users\Hooman Deghani\Python\Data Analysis\HTML\https___www.nerdwallet.com_article_mortgages_tips-for-first-time-home-buyers.txt",
r"C:\Users\Hooman Deghani\Python\Data Analysis\HTML\https___www.rocketmortgage.com_learn_first-time-home-buyer-tips.txt"
]

competitor_urls = [
"https://www.flexjobs.com/",
"https://www.hubspot.com/",
"https://www.investopedia.com/"
]

input_sq1 = r"C:\Users\Hooman Deghani\Python\Data Analysis\HTML\https___www.squareone.ca_resource-centres_home-buying-selling-moving_buying-a-home-for-the-first-time.txt"
sq1_url = "https://www.sq1.com/"

NLP(competitors_list, input_sq1, competitor_urls, sq1_url)